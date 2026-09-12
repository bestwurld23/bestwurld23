#!/usr/bin/env python3
"""Build XLSForm spreadsheets from the YAML sources in ``forms/src/``.

    python forms/build.py --pack packs/dev-fixture --out forms/dist

Why a build step at all: XLSForm's interchange format is a spreadsheet, but a
spreadsheet is not reviewable in a pull request and not editable without Excel.
The YAML is the source of truth; the ``.xlsx`` is a generated artifact, and the
site choice list is generated from the locale pack so a steward who adds a site
to the pack never hand-edits a form (PRD §7.1, §10).

Requires openpyxl:  pip install openpyxl
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

import yaml

# XLSForm column order. Only columns that some row actually uses are written, so
# a form with no skip logic does not ship an empty `relevant` column.
SURVEY_COLUMNS = [
    "type",
    "name",
    "label",
    "hint",
    "required",
    "relevant",
    "constraint",
    "constraint_message",
    "appearance",
    "default",
    "calculation",
    "choice_filter",
]
CHOICES_COLUMNS = ["list_name", "name", "label"]
SETTINGS_COLUMNS = [
    "form_title",
    "form_id",
    "version",
    "default_language",
    "instance_name",
    "style",
]


def load_sites(pack_dir: Path) -> list[dict[str, str]]:
    """Build the ``sites`` choice list from the pack's GeoJSON."""
    manifest = yaml.safe_load((pack_dir / "pack.yaml").read_text(encoding="utf-8"))
    sites_file = manifest.get("sites_file")
    if not sites_file:
        return []
    features = json.loads((pack_dir / sites_file).read_text(encoding="utf-8")).get("features", [])
    choices: list[dict[str, str]] = []
    for feature in features:
        properties = feature.get("properties", {})
        site_id = properties.get("id")
        if not site_id:
            continue
        label = properties.get("name", site_id)
        address = properties.get("address_text")
        choices.append(
            {
                "list_name": "sites",
                "name": site_id,
                "label": f"{label} — {address}" if address else str(label),
            }
        )
    choices.append({"list_name": "sites", "name": "other", "label": "Not in this list"})
    return choices


def build_rows(source: dict[str, Any], sites: list[dict[str, str]]) -> dict[str, list[dict]]:
    survey = [dict(row) for row in source["survey"]]

    choices: list[dict[str, str]] = []
    for list_name, entries in (source.get("choices") or {}).items():
        if entries == "from_pack":
            if list_name != "sites":
                raise ValueError(f"only the 'sites' list can come from the pack, not {list_name!r}")
            choices.extend(sites)
            continue
        for entry in entries:
            choices.append({"list_name": list_name, "name": entry["name"], "label": entry["label"]})

    settings = [source["settings"]]
    return {"survey": survey, "choices": choices, "settings": settings}


def used_columns(rows: list[dict], candidates: list[str]) -> list[str]:
    present = {key for row in rows for key, value in row.items() if value not in (None, "")}
    unknown = present - set(candidates)
    if unknown:
        raise ValueError(f"unknown XLSForm column(s): {sorted(unknown)}")
    return [column for column in candidates if column in present]


def write_xlsx(path: Path, sheets: dict[str, tuple[list[str], list[dict]]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:  # pragma: no cover - build-time dependency
        raise SystemExit("openpyxl is required to build XLSForms: pip install openpyxl") from None

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, (columns, rows) in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([_cell(row.get(column)) for column in columns])
        sheet.freeze_panes = "A2"
        for index, column in enumerate(columns, start=1):
            width = max([len(column)] + [len(_cell(r.get(column))) for r in rows] or [10])
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = min(
                max(width + 2, 12), 60
            )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    return str(value)


def build_form(source_path: Path, sites: list[dict[str, str]], out_dir: Path) -> Path:
    source = yaml.safe_load(source_path.read_text(encoding="utf-8"))
    rows = build_rows(source, sites)
    sheets = {
        "survey": (used_columns(rows["survey"], SURVEY_COLUMNS), rows["survey"]),
        "choices": (used_columns(rows["choices"], CHOICES_COLUMNS), rows["choices"]),
        "settings": (used_columns(rows["settings"], SETTINGS_COLUMNS), rows["settings"]),
    }
    out_path = out_dir / f"{source['settings']['form_id']}.xlsx"
    write_xlsx(out_path, sheets)
    return out_path


def main(argv: list[str] | None = None) -> int:
    repo_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pack", type=Path, default=repo_root / "packs" / "dev-fixture")
    parser.add_argument("--out", type=Path, default=repo_root / "forms" / "dist")
    parser.add_argument("--src", type=Path, default=repo_root / "forms" / "src")
    args = parser.parse_args(argv)

    sites = load_sites(args.pack)
    sources = sorted(args.src.glob("*.yaml"))
    if not sources:
        print(f"no form sources in {args.src}", file=sys.stderr)
        return 1
    for source_path in sources:
        out_path = build_form(source_path, sites, args.out)
        print(f"wrote {out_path.relative_to(repo_root)}")
    print(f"site choices from {args.pack.name}: {len(sites)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
