"""The XLSForm definitions (PRD §7.2).

The forms are configuration, but they are the part of the system a volunteer types
personal data into if we let them. So the rules are tested, not just documented:
no field may invite a name, and the built spreadsheet must match the sources and
the locale pack it was built from.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
import yaml

from cak_node.privacy import FORBIDDEN_COLUMNS, looks_like_pii

REPO_ROOT = Path(__file__).resolve().parents[1]
FORMS_SRC = REPO_ROOT / "forms" / "src"
EXPECTED_FORMS = {"cak_wait_time", "cak_turnaway", "cak_closure"}

sys.path.insert(0, str(REPO_ROOT / "forms"))


def form_sources() -> list[Path]:
    return sorted(FORMS_SRC.glob("*.yaml"))


def load(path: Path) -> dict:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def test_the_three_prd_forms_exist() -> None:
    ids = {load(p)["settings"]["form_id"] for p in form_sources()}
    assert ids == EXPECTED_FORMS


@pytest.mark.parametrize("source", form_sources(), ids=lambda p: p.stem)
def test_no_field_invites_personal_data(source: Path) -> None:
    """There is no name field to leave blank, because there is no name field."""
    form = load(source)
    names = {row["name"].lower() for row in form["survey"]}
    assert not names & FORBIDDEN_COLUMNS, sorted(names & FORBIDDEN_COLUMNS)
    for banned in ("name", "phone", "address", "email", "dob", "ssn", "id_number"):
        offenders = [n for n in names if banned in n and n not in {"site_other", "form_id"}]
        assert not offenders, f"{source.name} has field(s) {offenders} suggesting personal data"


@pytest.mark.parametrize("source", form_sources(), ids=lambda p: p.stem)
def test_free_text_fields_warn_against_names(source: Path) -> None:
    """A notes box without a warning is a notes box that collects names."""
    form = load(source)
    for row in form["survey"]:
        if row["type"] != "text" or row["name"].endswith("_url"):
            continue
        guidance = f"{row.get('hint', '')} {row.get('label', '')}".lower()
        assert "no name" in guidance or "not who" in guidance or "language only" in guidance, (
            f"{source.name}: free-text field {row['name']!r} does not warn against names"
        )


@pytest.mark.parametrize("source", form_sources(), ids=lambda p: p.stem)
def test_every_form_opens_with_the_rules(source: Path) -> None:
    form = load(source)
    intro = form["survey"][0]
    assert intro["type"] == "note"
    assert "name" in intro["label"].lower()


@pytest.mark.parametrize("source", form_sources(), ids=lambda p: p.stem)
def test_no_example_content_looks_like_real_personal_data(source: Path) -> None:
    text = source.read_text(encoding="utf-8")
    # Hints deliberately mention the words "name" and "address"; what must not
    # appear is anything matching an actual identifier.
    for line in text.splitlines():
        if looks_like_pii(line):
            pytest.fail(f"{source.name} contains something resembling personal data: {line!r}")


def test_turnaway_reason_is_not_a_legal_finding() -> None:
    """PRD §19: the reason code is what the person said, not a determination."""
    form = load(FORMS_SRC / "cak_turnaway.yaml")
    reason = next(r for r in form["survey"] if r["name"] == "reason_code")
    assert "not a legal decision" in reason["hint"].lower()
    assert "NOT A LEGAL FINDING" in form["survey"][0]["label"]

    codes = {c["name"] for c in form["choices"]["reason_code"]}
    assert {"id", "address", "closed", "language", "other"} <= codes


def test_closure_reports_are_marked_unofficial() -> None:
    form = load(FORMS_SRC / "cak_closure.yaml")
    assert "unofficial" in form["survey"][0]["label"].lower()
    new_location = next(r for r in form["survey"] if r["name"] == "new_location_text")
    assert "unofficial" in new_location["hint"].lower()


def test_wait_minutes_honour_the_prd_range() -> None:
    form = load(FORMS_SRC / "cak_wait_time.yaml")
    minutes = next(r for r in form["survey"] if r["name"] == "line_minutes")
    assert minutes["required"] is True
    assert ". >= 0 and . <= 480" in minutes["constraint"]


@pytest.mark.parametrize("source", form_sources(), ids=lambda p: p.stem)
def test_clock_drift_is_flagged_at_entry(source: Path) -> None:
    """PRD §7.2: flag a device clock that is obviously wrong."""
    form = load(source)
    observed = next(r for r in form["survey"] if r["name"] == "observed_at")
    assert observed["default"] == "now()"
    assert "now()" in observed["constraint"]
    assert observed["constraint_message"]


@pytest.mark.parametrize("source", form_sources(), ids=lambda p: p.stem)
def test_photo_fields_discourage_faces(source: Path) -> None:
    form = load(source)
    for row in form["survey"]:
        if row["type"] == "image":
            assert "face" in row.get("hint", "").lower(), row["name"]


def test_site_choices_are_generated_from_the_pack(pack_dir: Path) -> None:
    """Adding a site to a pack must not require editing a form."""
    import build as forms_build

    sites = forms_build.load_sites(pack_dir)
    ids = {choice["name"] for choice in sites}
    assert ids == {"site01", "site02", "site03", "other"}
    assert all(choice["list_name"] == "sites" for choice in sites)


def test_build_produces_importable_workbooks(pack_dir: Path, tmp_path: Path) -> None:
    pytest.importorskip("openpyxl")
    import build as forms_build
    from openpyxl import load_workbook

    sites = forms_build.load_sites(pack_dir)
    for source in form_sources():
        out = forms_build.build_form(source, sites, tmp_path)
        workbook = load_workbook(out)
        assert workbook.sheetnames == ["survey", "choices", "settings"]

        header = [cell.value for cell in workbook["survey"][1]]
        assert header[:3] == ["type", "name", "label"]

        settings_header = [cell.value for cell in workbook["settings"][1]]
        assert {"form_id", "form_title", "version"} <= set(settings_header)

        choice_lists = {
            row[0] for row in workbook["choices"].iter_rows(min_row=2, values_only=True)
        }
        referenced = {
            row["type"].split(" ", 1)[1]
            for row in load(source)["survey"]
            if row["type"].startswith("select_one ")
        }
        assert referenced <= choice_lists, f"{out.name} references an undefined choice list"


def test_committed_workbooks_match_the_sources(tmp_path: Path) -> None:
    """A stale dist/ is a form a steward imports that does not match review."""
    pytest.importorskip("openpyxl")
    import build as forms_build
    from openpyxl import load_workbook

    dist = REPO_ROOT / "forms" / "dist"
    if not dist.is_dir():
        pytest.skip("forms/dist not built")

    sites = forms_build.load_sites(REPO_ROOT / "packs" / "dev-fixture")
    for source in form_sources():
        form_id = load(source)["settings"]["form_id"]
        committed = dist / f"{form_id}.xlsx"
        if not committed.is_file():
            pytest.skip(f"{committed.name} not built")
        rebuilt = forms_build.build_form(source, sites, tmp_path)
        for sheet in ("survey", "choices", "settings"):
            stale = f"{committed.name} sheet {sheet!r} is stale; run python forms/build.py"
            assert _rows(load_workbook(committed)[sheet]) == _rows(load_workbook(rebuilt)[sheet]), (
                stale
            )


def _rows(sheet) -> list[tuple]:
    return [row for row in sheet.iter_rows(values_only=True)]
